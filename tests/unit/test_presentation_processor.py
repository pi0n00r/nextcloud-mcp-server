"""Presentations are read shape-by-shape, keeping slide/table structure."""

import io
import struct
import threading
import zlib

import anyio
import pytest
from pptx import Presentation
from pptx.util import Inches

from nextcloud_mcp_server.document_processors import presentation
from nextcloud_mcp_server.document_processors.base import ProcessorError
from nextcloud_mcp_server.document_processors.presentation import (
    PPTX_MIME,
    SLIDE_BOUNDARIES_KEY,
    PptxProcessor,
)

pytestmark = pytest.mark.unit


def _png_bytes(
    width: int, height: int, color: tuple[int, int, int] = (200, 50, 50)
) -> bytes:
    """A minimal valid PNG, built by hand so tests need no image library
    beyond what python-pptx itself already requires."""

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    raw = bytearray()
    for _ in range(height):
        raw.append(0)  # filter type: none
        raw.extend(bytes(color) * width)
    idat = zlib.compress(bytes(raw), 9)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", idat)
        + chunk(b"IEND", b"")
    )


def _deck(slides: list[dict]) -> bytes:
    """A .pptx built from ``[{"title": str, "body": [str, ...], "table": [[...]]}]``."""
    prs = Presentation()
    for spec in slides:
        if "table" in spec:
            slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
            rows = spec["table"]
            shape = slide.shapes.add_table(
                len(rows), len(rows[0]), Inches(1), Inches(1), Inches(4), Inches(2)
            )
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    shape.table.cell(r, c).text = value
            continue

        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if "title" in spec:
            slide.shapes.title.text = spec["title"]
        for line in spec.get("body", []):
            tf = slide.placeholders[1].text_frame
            para = tf.paragraphs[0] if not tf.text else tf.add_paragraph()
            para.text = line
        if "notes" in spec:
            slide.notes_slide.notes_text_frame.text = spec["notes"]

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_supported_mime_type_is_pptx_only():
    assert PptxProcessor().supported_mime_types == {PPTX_MIME}


async def test_text_frames_become_a_markdown_section():
    content = _deck([{"title": "Intro", "body": ["Line one", "Line two"]}])

    result = await PptxProcessor().process(content, PPTX_MIME, "d.pptx")

    assert "## Slide 1" in result.text
    assert "Intro" in result.text
    assert "Line one" in result.text
    assert "Line two" in result.text
    assert result.metadata["slide_count"] == 1
    assert result.metadata["parse_mode"] == "markdown"


async def test_a_table_shape_becomes_a_markdown_table():
    content = _deck([{"table": [["h1", "h2"], ["a", "b"]]}])

    result = await PptxProcessor().process(content, PPTX_MIME, "t.pptx")

    assert "| h1 | h2 |" in result.text
    assert "| a | b |" in result.text


async def test_pipe_in_a_cell_does_not_break_the_row():
    content = _deck([{"table": [["head"], ["a|b"]]}])

    result = await PptxProcessor().process(content, PPTX_MIME, "p.pptx")

    assert r"| a\|b |" in result.text


async def test_newline_in_a_cell_does_not_break_the_table():
    content = _deck([{"table": [["head"], ["line one\nline two"]]}])

    result = await PptxProcessor().process(content, PPTX_MIME, "n.pptx")

    assert "| line one line two |" in result.text


async def test_speaker_notes_are_included():
    content = _deck([{"title": "S", "body": ["x"], "notes": "Remember to smile"}])

    result = await PptxProcessor().process(content, PPTX_MIME, "notes.pptx")

    assert "**Notes:** Remember to smile" in result.text


async def test_each_slide_gets_its_own_boundary_span():
    content = _deck(
        [{"title": "First", "body": ["a"]}, {"title": "Second", "body": ["b"]}]
    )

    result = await PptxProcessor().process(content, PPTX_MIME, "two.pptx")

    spans = result.metadata[SLIDE_BOUNDARIES_KEY]
    assert [s["slide"] for s in spans] == [1, 2]
    # Offsets must index the returned text exactly -- they are what attributes a
    # chunk back to a slide, the presentation stand-in for a page number.
    for span in spans:
        segment = result.text[span["start_offset"] : span["end_offset"]]
        assert segment.startswith(f"## Slide {span['slide']}")
    assert spans[0]["end_offset"] == spans[1]["start_offset"]


async def test_empty_slide_is_skipped_not_emitted_as_an_empty_section():
    content = _deck([{}, {"title": "Real", "body": ["x"]}])

    result = await PptxProcessor().process(content, PPTX_MIME, "e.pptx")

    assert "## Slide 1" not in result.text
    assert "## Slide 2" in result.text
    assert [s["slide"] for s in result.metadata[SLIDE_BOUNDARIES_KEY]] == [2]


async def test_health_check_is_true_once_pptx_is_importable():
    assert await PptxProcessor().health_check() is True


def _save(prs) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def test_notes_page_without_a_notes_placeholder_does_not_fail_the_deck():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Still readable"
    notes = slide.notes_slide
    for placeholder in list(notes.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    assert notes.notes_text_frame is None

    result = await PptxProcessor().process(_save(prs), PPTX_MIME, "nn.pptx")

    assert "Still readable" in result.text
    assert "**Notes:**" not in result.text


async def test_text_inside_nested_group_shapes_is_extracted():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank: groups only
    outer = slide.shapes.add_group_shape()
    outer.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = "outer text"
    inner = outer.shapes.add_group_shape()
    inner.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = "inner text"

    result = await PptxProcessor().process(_save(prs), PPTX_MIME, "g.pptx")

    assert "## Slide 1" in result.text
    assert "outer text" in result.text
    assert "inner text" in result.text


async def test_deck_inflating_past_the_cap_is_refused(monkeypatch):
    content = _deck([{"title": "T", "body": ["x"]}])
    monkeypatch.setattr(presentation, "MAX_UNCOMPRESSED_BYTES", len(content))
    processor = PptxProcessor()

    with pytest.raises(ProcessorError, match="uncompressed size"):
        await processor.process(content, PPTX_MIME, "bomb.pptx")


async def test_a_stuck_parse_is_abandoned_when_the_caller_times_out(monkeypatch):
    release = threading.Event()
    monkeypatch.setattr(presentation, "_extract_deck", lambda _c: release.wait(5))
    processor = PptxProcessor()

    try:
        with pytest.raises(TimeoutError), anyio.fail_after(0.2):
            await processor.process(b"", PPTX_MIME, "slow.pptx")
    finally:
        release.set()


# --- picture eligibility (ADR-037) -------------------------------------------


def test_eligible_picture_skips_unsupported_content_type(mocker):
    shape = mocker.Mock()
    shape.image.content_type = (
        "image/x-emf"  # vector paste, docling can't read it either
    )
    shape.image.size = (500, 500)

    assert presentation._eligible_picture(shape) is None


def test_eligible_picture_skips_pictures_below_the_size_floor(mocker):
    shape = mocker.Mock()
    shape.image.content_type = "image/png"
    shape.image.size = (40, 40)  # below MIN_CAPTION_PICTURE_PX -- logo/icon-sized

    assert presentation._eligible_picture(shape) is None


def test_eligible_picture_accepts_a_supported_large_picture(mocker):
    shape = mocker.Mock()
    shape.image.content_type = "image/png"
    shape.image.size = (200, 150)
    shape.image.blob = b"raw-bytes"

    picture = presentation._eligible_picture(shape)

    assert picture is not None
    assert picture.content_type == "image/png"
    assert picture.blob == b"raw-bytes"


def test_eligible_picture_returns_none_for_an_unreadable_image_part(mocker):
    shape = mocker.Mock()
    type(shape).image = mocker.PropertyMock(side_effect=Exception("corrupt part"))

    assert presentation._eligible_picture(shape) is None


# --- picture captioning end to end (ADR-037) ---------------------------------


def _deck_with_picture(image_bytes: bytes) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    slide.shapes.add_picture(
        io.BytesIO(image_bytes), Inches(1), Inches(1), Inches(2), Inches(2)
    )
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def test_picture_is_counted_but_not_captioned_when_disabled():
    content = _deck_with_picture(_png_bytes(120, 120))

    result = await PptxProcessor().process(content, PPTX_MIME, "pic.pptx")

    assert result.metadata["pptx_pictures_found"] == 1
    assert "pptx_pictures_captioned" not in result.metadata
    assert "*Image:" not in result.text


async def test_small_picture_is_not_counted_as_eligible():
    content = _deck_with_picture(_png_bytes(40, 40))

    result = await PptxProcessor().process(content, PPTX_MIME, "small.pptx")

    assert result.metadata["pptx_pictures_found"] == 0


async def test_caption_images_flag_without_docling_url_stays_disabled(
    mocker, monkeypatch
):
    convert = mocker.AsyncMock()
    monkeypatch.setattr(presentation, "convert_file", convert)
    content = _deck_with_picture(_png_bytes(120, 120))

    processor = PptxProcessor(caption_images=True, docling_api_url=None)
    result = await processor.process(content, PPTX_MIME, "pic.pptx")

    convert.assert_not_called()
    assert "pptx_pictures_captioned" not in result.metadata


async def test_picture_is_captioned_when_docling_is_configured(mocker, monkeypatch):
    convert = mocker.AsyncMock(return_value={"md_content": "a red square"})
    monkeypatch.setattr(presentation, "convert_file", convert)
    content = _deck_with_picture(_png_bytes(120, 120))

    processor = PptxProcessor(
        caption_images=True, docling_api_url="https://docling:5001"
    )
    result = await processor.process(content, PPTX_MIME, "pic.pptx")

    assert "*Image: a red square*" in result.text
    assert result.metadata["pptx_pictures_found"] == 1
    assert result.metadata["pptx_pictures_captioned"] == 1
    # The picture's own bytes/content type are forwarded, not re-derived.
    args, kwargs = convert.call_args
    assert args[0] == "https://docling:5001"
    assert args[2] == "image/png"
    assert kwargs["to_formats"] == ["md"]


async def test_caption_failure_does_not_fail_the_deck(mocker, monkeypatch):
    convert = mocker.AsyncMock(side_effect=ProcessorError("docling unreachable"))
    monkeypatch.setattr(presentation, "convert_file", convert)
    content = _deck_with_picture(_png_bytes(120, 120))

    processor = PptxProcessor(
        caption_images=True, docling_api_url="https://docling:5001"
    )
    result = await processor.process(content, PPTX_MIME, "pic.pptx")

    assert result.success is True
    assert "*Image:" not in result.text
    assert result.metadata["pptx_pictures_found"] == 1
    assert result.metadata["pptx_pictures_captioned"] == 0


async def test_caption_max_images_caps_docling_round_trips(mocker, monkeypatch):
    convert = mocker.AsyncMock(return_value={"md_content": "caption"})
    monkeypatch.setattr(presentation, "convert_file", convert)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for _ in range(3):
        slide.shapes.add_picture(
            io.BytesIO(_png_bytes(120, 120)), Inches(1), Inches(1), Inches(2), Inches(2)
        )
    buf = io.BytesIO()
    prs.save(buf)

    processor = PptxProcessor(
        caption_images=True,
        docling_api_url="https://docling:5001",
        caption_max_images=2,
    )
    result = await processor.process(buf.getvalue(), PPTX_MIME, "many.pptx")

    assert convert.await_count == 2
    assert result.metadata["pptx_pictures_found"] == 3
    assert result.metadata["pptx_pictures_captioned"] == 2
