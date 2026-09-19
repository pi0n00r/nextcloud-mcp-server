"""Word documents are read paragraph-by-paragraph into markdown (ADR-038)."""

import io

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from nextcloud_mcp_server.document_processors import _ooxml
from nextcloud_mcp_server.document_processors._ooxml import PictureCaptioner
from nextcloud_mcp_server.document_processors.base import ProcessorError
from nextcloud_mcp_server.document_processors.word import DOCX_MIME, DocxProcessor

pytestmark = pytest.mark.unit

DOCLING = "https://docling:5001"


def _png(width: int, height: int) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), (200, 50, 50)).save(buf, "PNG")
    return buf.getvalue()


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _captioning(**kwargs) -> DocxProcessor:
    return DocxProcessor(
        captioner=PictureCaptioner(
            caption_images=True, docling_api_url=DOCLING, **kwargs
        )
    )


def test_supported_mime_type_is_docx_only():
    assert DocxProcessor().supported_mime_types == {DOCX_MIME}


async def test_headings_lists_and_paragraphs_become_markdown():
    doc = Document()
    doc.add_heading("Report", level=0)
    doc.add_heading("Scope", level=1)
    doc.add_heading("Detail", level=2)
    doc.add_paragraph("Plain prose.")
    doc.add_paragraph("first item", style="List Bullet")
    doc.add_paragraph("")  # empty paragraphs are dropped

    result = await DocxProcessor().process(_save(doc), DOCX_MIME, "d.docx")

    assert result.text == (
        "# Report\n\n# Scope\n\n## Detail\n\nPlain prose.\n\n- first item\n"
    )
    assert result.metadata["parse_mode"] == "markdown"
    assert result.metadata["pictures_found"] == 0


async def test_direct_numbering_on_a_plain_style_is_a_list_item():
    doc = Document()
    paragraph = doc.add_paragraph("numbered without a List style")
    paragraph._p.get_or_add_pPr().get_or_add_numPr()

    result = await DocxProcessor().process(_save(doc), DOCX_MIME, "n.docx")

    assert result.text == "- numbered without a List style\n"


async def test_a_table_becomes_a_markdown_table_in_body_order():
    doc = Document()
    doc.add_paragraph("before")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate([["h1", "h2"], ["a|b", "line one\nline two"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.add_paragraph("after")

    result = await DocxProcessor().process(_save(doc), DOCX_MIME, "t.docx")

    assert result.text == (
        "before\n\n| h1 | h2 |\n| --- | --- |\n| a\\|b | line one line two |\n\nafter\n"
    )


async def test_empty_document_yields_empty_text():
    result = await DocxProcessor().process(_save(Document()), DOCX_MIME, "e.docx")

    assert result.text == ""


async def test_document_inflating_past_the_cap_is_refused(monkeypatch):
    content = _save(Document())
    monkeypatch.setattr(_ooxml, "MAX_UNCOMPRESSED_BYTES", 1)

    with pytest.raises(ProcessorError, match="uncompressed size"):
        await DocxProcessor().process(content, DOCX_MIME, "bomb.docx")


def _doc_with_pictures(*sizes: tuple[int, int]) -> bytes:
    doc = Document()
    doc.add_paragraph("intro")
    for width, height in sizes:
        doc.add_picture(io.BytesIO(_png(width, height)), width=Inches(1))
    doc.add_paragraph("outro")
    return _save(doc)


async def test_picture_is_counted_but_not_captioned_when_disabled(mocker, monkeypatch):
    convert = mocker.AsyncMock()
    monkeypatch.setattr(_ooxml, "convert_file", convert)

    result = await DocxProcessor().process(
        _doc_with_pictures((120, 120), (40, 40)), DOCX_MIME, "p.docx"
    )

    convert.assert_not_called()
    # The 40px picture is decorative and never counted.
    assert result.metadata["pictures_found"] == 1
    assert "pictures_captioned" not in result.metadata
    assert "*Image:" not in result.text


async def test_caption_lands_after_the_paragraph_holding_the_picture(
    mocker, monkeypatch
):
    convert = mocker.AsyncMock(return_value={"md_content": "a red square"})
    monkeypatch.setattr(_ooxml, "convert_file", convert)

    result = await _captioning().process(
        _doc_with_pictures((120, 120)), DOCX_MIME, "p.docx"
    )

    assert result.text == "intro\n\n*Image: a red square*\n\noutro\n"
    assert result.metadata["pictures_captioned"] == 1
    args, _ = convert.call_args
    assert args[2] == "image/png"


async def test_caption_failure_does_not_fail_the_document(mocker, monkeypatch):
    convert = mocker.AsyncMock(side_effect=ProcessorError("docling unreachable"))
    monkeypatch.setattr(_ooxml, "convert_file", convert)

    result = await _captioning().process(
        _doc_with_pictures((120, 120)), DOCX_MIME, "p.docx"
    )

    assert result.text == "intro\n\noutro\n"
    assert result.metadata["pictures_found"] == 1
    assert result.metadata["pictures_captioned"] == 0


async def test_caption_max_images_caps_docling_round_trips(mocker, monkeypatch):
    convert = mocker.AsyncMock(
        side_effect=[{"md_content": "one"}, {"md_content": "two"}]
    )
    monkeypatch.setattr(_ooxml, "convert_file", convert)

    result = await _captioning(caption_max_images=2).process(
        _doc_with_pictures((120, 120), (120, 120), (120, 120)), DOCX_MIME, "m.docx"
    )

    assert convert.await_count == 2
    # Captions keep document order even with several pictures in a row.
    assert result.text == "intro\n\n*Image: one*\n\n*Image: two*\n\noutro\n"
    assert result.metadata["pictures_found"] == 3
    assert result.metadata["pictures_captioned"] == 2


async def test_health_check_is_true_once_docx_is_importable():
    assert await DocxProcessor().health_check() is True


async def test_two_pictures_in_one_paragraph_keep_their_order(mocker, monkeypatch):
    convert = mocker.AsyncMock(
        side_effect=[{"md_content": "first"}, {"md_content": "second"}]
    )
    monkeypatch.setattr(_ooxml, "convert_file", convert)
    doc = Document()
    doc.add_paragraph("intro")
    paragraph = doc.add_paragraph("")
    for _ in range(2):
        paragraph.add_run().add_picture(io.BytesIO(_png(120, 120)), width=Inches(1))
    doc.add_paragraph("outro")

    result = await _captioning().process(_save(doc), DOCX_MIME, "two.docx")

    assert result.text == "intro\n\n*Image: first*\n\n*Image: second*\n\noutro\n"
