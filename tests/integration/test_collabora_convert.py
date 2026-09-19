"""Legacy and ODF office files read through a live Collabora Online (ADR-039).

Gated on ``COLLABORA_URL``. Run the stack with the "collabora" profile and give
the single-user MCP service ``COLLABORA_URL=http://collabora:9980``; set
``COLLABORA_URL`` in this process too (any non-empty value) so the skipif lets
the tests run. The CI "collabora" lane does both and selects ``-m collabora``.

The legacy fixtures are made by the same coolwsd container this suite tests
(published on localhost:9980): there is no pure-Python writer for .doc/.xls/.ppt.
A .doc is made via .odt because LibreOffice's direct .docx -> .doc export drops
python-docx's tables -- a fixture artifact, not something a real Word .doc hits.
"""

import json
import os
import uuid
from io import BytesIO

import httpx
import pytest
from docx import Document
from mcp.client.session import ClientSession
from openpyxl import Workbook
from pptx import Presentation

from nextcloud_mcp_server.client import NextcloudClient

pytestmark = [
    pytest.mark.integration,
    pytest.mark.collabora,
    pytest.mark.skipif(
        not os.getenv("COLLABORA_URL"), reason="needs the collabora profile"
    ),
]

COOLWSD = "http://localhost:9980/cool/convert-to"
MARKER = "CollaboraConvertMarker"


def _convert(data: bytes, filename: str, target: str) -> bytes:
    response = httpx.post(
        f"{COOLWSD}/{target}", files={"data": (filename, data)}, timeout=120
    )
    response.raise_for_status()
    return response.content


def _save(obj) -> bytes:
    buf = BytesIO()
    obj.save(buf)
    return buf.getvalue()


def _docx_with_table() -> bytes:
    doc = Document()
    doc.add_heading(f"Report {MARKER}", level=1)
    table = doc.add_table(rows=2, cols=2, style="Table Grid")
    for (r, c), value in zip(
        [(0, 0), (0, 1), (1, 0), (1, 1)], ["Qtr", "Rev", "Q1", "42"]
    ):
        table.cell(r, c).text = value
    return _save(doc)


def _xlsx() -> bytes:
    wb = Workbook()
    wb.active.title = "Revenue"
    wb.active.append(["Quarter", MARKER])
    wb.active.append(["Q1", 42])
    return _save(wb)


def _pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = f"Slide {MARKER}"
    return _save(prs)


def _fixtures() -> dict[str, tuple[bytes, str]]:
    odt = _convert(_docx_with_table(), "t.docx", "odt")
    return {
        "legacy.doc": (_convert(odt, "t.odt", "doc"), "application/msword"),
        "legacy.xls": (_convert(_xlsx(), "t.xlsx", "xls"), "application/vnd.ms-excel"),
        "legacy.ppt": (
            _convert(_pptx(), "t.pptx", "ppt"),
            "application/vnd.ms-powerpoint",
        ),
        "open.odt": (odt, "application/vnd.oasis.opendocument.text"),
    }


@pytest.fixture
async def uploaded(nc_client: NextcloudClient):
    test_dir = f"mcp_collabora_{uuid.uuid4().hex[:8]}"
    await nc_client.webdav.create_directory(test_dir)
    for name, (data, mime) in _fixtures().items():
        await nc_client.webdav.write_file(f"{test_dir}/{name}", data, content_type=mime)
    yield test_dir
    try:
        await nc_client.webdav.delete_resource(test_dir)
    except Exception:
        pass  # Ignore cleanup errors


async def _read(nc_mcp_client: ClientSession, path: str) -> dict:
    result = await nc_mcp_client.call_tool(
        "nc_webdav_read_file", arguments={"path": path}
    )
    content = result.content[0]
    return json.loads(content.text if hasattr(content, "text") else str(content))


@pytest.mark.parametrize(
    ("name", "processor", "expected"),
    [
        ("legacy.doc", "collabora+word", ["# Report", "| Q1 | 42 |"]),
        ("legacy.xls", "collabora+spreadsheet", ["## Sheet: Revenue", "| Q1 | 42 |"]),
        ("legacy.ppt", "collabora+presentation", ["## Slide 1"]),
        ("open.odt", "collabora+word", ["# Report", "| Q1 | 42 |"]),
    ],
)
async def test_converted_file_reads_as_native_markdown(
    nc_mcp_client: ClientSession, uploaded: str, name, processor, expected
):
    result = await _read(nc_mcp_client, f"{uploaded}/{name}")

    assert result["parse_status"] == "parsed", result.get("parse_notes")
    assert result["parse_processor"] == processor
    assert MARKER in result["content"]
    for fragment in expected:
        assert fragment in result["content"]
