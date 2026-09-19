"""End-to-end coverage of reading a document through nc_webdav_read_file (Deck #894).

Runs against the single-user MCP service (port 8000) with the built-in PDF tiers
-- no optional processor, no OCR backend, nothing enabled beyond a default
deployment. That is the point: an agent reading a PDF gets its text out of the
box, and the response says which tier produced it.
"""

import json
import logging
import uuid
from io import BytesIO

import pytest
from mcp.client.session import ClientSession
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from nextcloud_mcp_server.client import NextcloudClient

logger = logging.getLogger(__name__)

pytestmark = pytest.mark.integration

MARKER = "ParseDocumentIntegrationMarker"


@pytest.fixture
async def test_base_path(nc_client: NextcloudClient):
    test_dir = f"mcp_test_doc_read_{uuid.uuid4().hex[:8]}"
    await nc_client.webdav.create_directory(test_dir)
    yield test_dir
    try:
        await nc_client.webdav.delete_resource(test_dir)
    except Exception:
        pass  # Ignore cleanup errors


#: Ordinary prose, on purpose. The tier-0 classifier scores the extracted text
#: layer, and a page holding one short unspaced token reads as low-quality: it
#: escalates fast->structured (verified — a single ``drawString(MARKER)`` page
#: lands on the structured tier and comes back as markdown). That would make this
#: file a test of classifier heuristics rather than of the read path, and would
#: let the markdown test below pass without markdown ever being requested. Keep
#: the body realistic if you touch it.
_BODY = [
    "This document exists so an integration test can read a real PDF back",
    "through the MCP tool and check that its text comes out as text rather",
    "than as base64, which is what this tool used to return for any document.",
    "The wording is deliberately ordinary: several lines of normal words give",
    "the tier-0 classifier a text layer it can be confident about, so the read",
    "stops at the cheap fast tier instead of escalating to recover it.",
]


@pytest.fixture
async def text_layer_pdf(nc_client: NextcloudClient, test_base_path: str):
    """A born-digital single-page PDF with a healthy text layer."""
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    c.drawString(72, 750, f"Quarterly Report {MARKER}")
    for offset, line in enumerate(_BODY, start=1):
        c.drawString(72, 750 - offset * 20, line)
    c.save()

    path = f"{test_base_path}/document.pdf"
    await nc_client.webdav.write_file(
        path, buffer.getvalue(), content_type="application/pdf"
    )
    return path


def _read_result(mcp_result) -> dict:
    content = mcp_result.content[0]
    text = content.text if hasattr(content, "text") else str(content)
    return json.loads(text)


async def _read(nc_mcp_client: ClientSession, path: str, **arguments) -> dict:
    return _read_result(
        await nc_mcp_client.call_tool(
            "nc_webdav_read_file", arguments={"path": path, **arguments}
        )
    )


async def test_pdf_reads_as_text_by_default(
    nc_mcp_client: ClientSession, text_layer_pdf: str
):
    """The default is text, not base64 -- the whole point of the card."""
    result = await _read(nc_mcp_client, text_layer_pdf)

    assert result["parsed"] is True
    assert result["parse_status"] == "parsed"
    assert result["parse_tier"] == "fast"
    assert result["content_format"] == "text"
    assert result.get("encoding") is None
    assert MARKER in result["content"]
    # A clean parse has nothing to disclose.
    assert result["parse_notes"] == []
    # The etag survives the streamed download, so the caller can still write back
    # with an If-Match precondition.
    assert result["etag"]


async def test_markdown_mode_reconstructs_structure(
    nc_mcp_client: ClientSession, text_layer_pdf: str
):
    """Asking for markdown promotes the read to the structured tier."""
    result = await _read(nc_mcp_client, text_layer_pdf, parse_document="markdown")

    assert result["parse_status"] == "parsed"
    assert result["parse_tier"] == "structured"
    assert result["content_format"] == "markdown"
    assert MARKER in result["content"]


async def test_raw_mode_returns_the_bytes(
    nc_mcp_client: ClientSession, text_layer_pdf: str
):
    """The raw path stays reachable for a caller that wants the file itself."""
    result = await _read(nc_mcp_client, text_layer_pdf, parse_document="raw")

    assert result["parsed"] is False
    assert result["parse_status"] == "skipped"
    assert result["encoding"] == "base64"
    assert result["content_format"] == "base64"
    assert result["parse_notes"] == []


async def test_text_file_is_untouched_by_the_parse_argument(
    nc_client: NextcloudClient, nc_mcp_client: ClientSession, test_base_path: str
):
    """A plain text file needs no processor and is reported as such."""
    path = f"{test_base_path}/notes.txt"
    await nc_client.webdav.write_file(path, b"hello world", content_type="text/plain")

    result = await _read(nc_mcp_client, path)

    assert result["content"] == "hello world"
    assert result["parse_status"] == "not_applicable"
    assert result["parsed"] is False
    assert result["parse_notes"] == []


async def test_read_returns_a_link_that_opens_the_file(
    nc_client: NextcloudClient, nc_mcp_client: ClientSession, text_layer_pdf: str
):
    """The response carries a Files link, so a caller can offer the original.

    Pinned end-to-end because the fileid it needs is NOT in the read itself: a
    WebDAV GET returns no ``OC-FileId`` (only a PUT does), so the tool resolves
    it with a separate PROPFIND. A unit test on the URL builder cannot show that
    lookup actually happens, or that it happens on the parsed path too.
    """
    result = await _read(nc_mcp_client, text_layer_pdf)

    file_id = await nc_client.webdav.get_fileid(text_layer_pdf)
    assert file_id, "fixture file has no fileid; the assertion below is vacuous"
    assert result["url"], (
        "no file link on the read; the mcp service resolves "
        "nextcloud_browser_url from NEXTCLOUD_PUBLIC_ISSUER_URL, so an empty "
        "url means the builder was not reached"
    )
    assert result["url"].endswith(f"/index.php/f/{file_id}")


async def test_raw_read_is_linked_too(
    nc_client: NextcloudClient, nc_mcp_client: ClientSession, text_layer_pdf: str
):
    """The raw path builds its response at a different site than the parsed one.

    Five sites build a ``ReadFileResponse``; the link is stamped once, after the
    fact, precisely so a second one cannot silently return None.
    """
    result = await _read(nc_mcp_client, text_layer_pdf, parse_document="raw")

    file_id = await nc_client.webdav.get_fileid(text_layer_pdf)
    assert result["url"].endswith(f"/index.php/f/{file_id}")


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


async def test_docx_reads_as_markdown_without_an_optional_processor(
    nc_client: NextcloudClient, nc_mcp_client: ClientSession, test_base_path: str
):
    """The native reader (ADR-038) needs no unstructured service. A picture is
    counted even with captioning off, and the response says so."""
    from docx import Document  # noqa: PLC0415
    from docx.shared import Inches  # noqa: PLC0415
    from PIL import Image  # noqa: PLC0415

    doc = Document()
    doc.add_heading(f"Report {MARKER}", level=1)
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate([["Quarter", "Revenue"], ["Q1", "42"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    image = BytesIO()
    Image.new("RGB", (120, 120), (200, 50, 50)).save(image, "PNG")
    doc.add_picture(BytesIO(image.getvalue()), width=Inches(1))
    buffer = BytesIO()
    doc.save(buffer)
    path = f"{test_base_path}/document.docx"
    await nc_client.webdav.write_file(path, buffer.getvalue(), content_type=DOCX_MIME)

    result = await _read(nc_mcp_client, path)

    assert result["parse_status"] == "parsed"
    assert f"# Report {MARKER}" in result["content"]
    assert "| Q1 | 42 |" in result["content"]
    assert any("1 picture(s)" in note for note in result["parse_notes"])


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def test_xlsx_reads_as_markdown_tables_without_an_optional_processor(
    nc_client: NextcloudClient, nc_mcp_client: ClientSession, test_base_path: str
):
    """The native reader (ADR-038) renders one markdown table per sheet."""
    from openpyxl import Workbook  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Quarter", MARKER])
    ws.append(["Q1", 42])
    buffer = BytesIO()
    wb.save(buffer)
    path = f"{test_base_path}/workbook.xlsx"
    await nc_client.webdav.write_file(path, buffer.getvalue(), content_type=XLSX_MIME)

    result = await _read(nc_mcp_client, path)

    assert result["parse_status"] == "parsed"
    assert "## Sheet: Revenue" in result["content"]
    assert f"| Quarter | {MARKER} |" in result["content"]
    assert "| Q1 | 42 |" in result["content"]
    assert result["parse_notes"] == []
