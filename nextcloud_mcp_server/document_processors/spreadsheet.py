"""Spreadsheets, read cell-by-cell rather than through a PDF rendition.

``.xlsx`` is OOXML, so ``openpyxl`` reads it directly (ADR-038). Rendering to
PDF would honour the print layout and drop whatever falls outside it -- in #1265
a real workbook rendered to 567 cells against the 2013 a direct read recovers.
Legacy ``.xls`` (OLE2) is out of scope: openpyxl cannot open it.

Optionally (``OFFICE_CAPTION_IMAGES`` + ``DOCLING_API_URL``, ADR-037), pictures
embedded in the workbook are captioned by docling-serve and listed in a
trailing ``## Images`` section.
"""

import io
import logging
import mimetypes
import re
import zipfile
from collections.abc import Awaitable, Callable
from typing import Any, Optional

from anyio.to_thread import run_sync

from ._ooxml import (
    Picture,
    PictureCaptioner,
    caption_line,
    check_zip_size,
    escape_cell,
    picture_if_eligible,
    picture_metadata,
    render_table,
)
from .base import DocumentProcessor, ProcessingResult, ProcessorError

logger = logging.getLogger(__name__)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class XlsxProcessor(DocumentProcessor):
    """Extract ``.xlsx`` as one markdown table per sheet."""

    def __init__(self, *, captioner: PictureCaptioner | None = None) -> None:
        self._captioner = captioner or PictureCaptioner()

    @property
    def name(self) -> str:
        return "spreadsheet"

    @property
    def tier(self) -> str:
        return "fast"

    @property
    def supported_mime_types(self) -> set[str]:
        return {XLSX_MIME}

    async def process(
        self,
        content: bytes,
        content_type: str,
        filename: Optional[str] = None,
        options: Optional[dict[str, Any]] = None,
        progress_callback: Optional[
            Callable[[float, Optional[float], Optional[str]], Awaitable[None]]
        ] = None,
    ) -> ProcessingResult:
        try:
            sections, sheet_names, pictures = await run_sync(
                _extract_workbook, content, abandon_on_cancel=True
            )
        except Exception as exc:
            raise ProcessorError(f"Spreadsheet parse failed: {exc}") from exc

        captions = (
            await self._captioner.caption_all(pictures)
            if self._captioner.enabled
            else None
        )
        # ponytail: captions are not tied to a sheet; resolve the
        # xl/drawings rels if per-sheet placement turns out to matter.
        lines = [caption_line(c) for c in captions or [] if c]
        if lines:
            sections.append("## Images\n\n" + "\n\n".join(lines))

        text = "\n\n".join(sections) + ("\n" if sections else "")
        return ProcessingResult(
            text=text,
            metadata={
                "sheet_count": len(sheet_names),
                "sheet_names": sheet_names,
                "text_length": len(text),
                "parse_mode": "markdown",
                **picture_metadata(len(pictures), captions),
            },
            processor=self.name,
            success=True,
        )

    async def health_check(self) -> bool:
        try:
            import openpyxl  # noqa: F401, PLC0415
        except ImportError:
            return False
        return True


def _sheet_table(sheet: Any) -> str:
    """One sheet as a markdown table, or "" when it holds nothing.

    Empty rows and trailing empty cells are dropped, and rows padded to the
    widest. The first non-empty row becomes the header -- a guess, since a
    sheet may open with a title banner, but every cell survives either way.
    """
    # A workbook's stored dimension can be stale, and read-only mode trusts
    # it; resetting makes iter_rows read to the real end of the sheet.
    sheet.reset_dimensions()
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        cells = [escape_cell("" if v is None else str(v)) for v in row]
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            rows.append(cells)
    width = max((len(r) for r in rows), default=0)
    return render_table([r + [""] * (width - len(r)) for r in rows])


def _natural_key(name: str) -> list[int | str]:
    """Sort ``image2.png`` before ``image10.png``: the caption cap should take
    the first pictures in insertion order, not in lexicographic order."""
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", name)]


def _media_pictures(content: bytes) -> list[Picture]:
    """Eligible pictures stored in the package's ``xl/media/`` folder."""
    from PIL import Image  # noqa: PLC0415 -- keep the import off the hot path

    pictures = []
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        for name in sorted(zf.namelist(), key=_natural_key):
            if not name.startswith("xl/media/"):
                continue
            content_type = mimetypes.guess_type(name)[0] or ""
            blob = zf.read(name)
            try:
                with Image.open(io.BytesIO(blob)) as image:
                    size = image.size  # header only, no full decode
            except Exception as exc:  # noqa: BLE001
                # An unreadable image (or an EMF Pillow cannot size) costs
                # this picture only.
                logger.debug("Skipping unreadable xlsx picture %s: %s", name, exc)
                continue
            picture = picture_if_eligible(content_type, size, blob)
            if picture is not None:
                pictures.append(picture)
    return pictures


def _extract_workbook(content: bytes) -> tuple[list[str], list[str], list[Picture]]:
    """Per-sheet markdown sections, the sheet names, and eligible pictures.
    Runs in a worker thread; captioning happens afterwards, in ``process()``."""
    import openpyxl  # noqa: PLC0415 -- keep the import off the hot path

    check_zip_size(content)
    # data_only: the cached result of a formula, not "=SUM(A1:A9)" -- the value
    # is what a reader searches for. read_only streams rows instead of building
    # the whole object graph, so a large workbook stays bounded (it also skips
    # images, which _media_pictures reads straight from the package instead).
    workbook = openpyxl.load_workbook(
        io.BytesIO(content), data_only=True, read_only=True
    )
    try:
        sections = []
        for sheet in workbook.worksheets:
            table = _sheet_table(sheet)
            if table:
                sections.append(f"## Sheet: {sheet.title}\n\n{table}")
        sheet_names = list(workbook.sheetnames)
    finally:
        workbook.close()
    return sections, sheet_names, _media_pictures(content)
