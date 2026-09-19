"""Presentations, read shape-by-shape rather than through a PDF rendition.

``.pptx`` is already OOXML (a zip of XML parts), so ``python-pptx`` reads it
directly -- no LibreOffice/``soffice`` dependency. Legacy ``.ppt`` (OLE2) is
out of scope: python-pptx cannot open the OLE2 container, only the OOXML one.

Optionally (``OFFICE_CAPTION_IMAGES`` + ``DOCLING_API_URL``, see ADR-037), a
raster picture pasted onto a slide is captioned by sending it to the same
docling-serve instance the images-only ``DoclingProcessor`` uses, and the
caption is appended to that slide's text. This only reaches actual pictures --
native vector diagrams (SmartArt, freeform/connector shapes) have no image
part to send and are unaffected; python-pptx has no rendering engine to turn
those into pixels.
"""

import io
import logging
from collections.abc import Awaitable, Callable
from dataclasses import dataclass, field
from typing import Any, Optional

from anyio.to_thread import run_sync

from ._ooxml import (
    Picture,
    PictureCaptioner,
    caption_line,
    check_zip_size,
    escape_cell,
    picture_if_eligible,
    picture_metadata,
    render_table,
)
from .base import DocumentProcessor, ProcessingResult, ProcessorError

logger = logging.getLogger(__name__)

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# Per-slide spans into the joined text: {"slide", "start_offset", "end_offset"}.
# The presentation counterpart to a PDF's ``page_boundaries`` -- lets a chunk be
# attributed to the slide it came from.
SLIDE_BOUNDARIES_KEY = "slide_boundaries"


@dataclass
class _SlideData:
    """One slide's extracted blocks, notes and picture shapes, pre-captioning."""

    index: int
    blocks: list[str]
    pictures: list[Picture] = field(default_factory=list)
    notes: Optional[str] = None


class PptxProcessor(DocumentProcessor):
    """Extract ``.pptx`` as one markdown section per slide."""

    def __init__(self, *, captioner: PictureCaptioner | None = None) -> None:
        # One captioner is shared by every OOXML reader (see __init__.py);
        # the default one is disabled.
        self._captioner = captioner or PictureCaptioner()

    @property
    def name(self) -> str:
        return "presentation"

    @property
    def tier(self) -> str:
        return "fast"

    @property
    def supported_mime_types(self) -> set[str]:
        return {PPTX_MIME}

    async def _caption_all_pictures(
        self, slides: list["_SlideData"]
    ) -> list[str | None]:
        """Caption every slide's pictures (capped), appending each success to
        its slide's blocks. Returns the per-picture captions."""
        pictures = [p for slide in slides for p in slide.pictures]
        captions = await self._captioner.caption_all(pictures)
        it = iter(captions)
        for slide in slides:
            for _ in slide.pictures:
                if caption := next(it):
                    slide.blocks.append(caption_line(caption))
        return captions

    async def process(
        self,
        content: bytes,
        content_type: str,
        filename: Optional[str] = None,
        options: Optional[dict[str, Any]] = None,
        progress_callback: Optional[
            Callable[[float, Optional[float], Optional[str]], Awaitable[None]]
        ] = None,
    ) -> ProcessingResult:
        try:
            slides, slide_count = await run_sync(
                _extract_deck, content, abandon_on_cancel=True
            )
        except Exception as exc:
            raise ProcessorError(f"Presentation parse failed: {exc}") from exc

        pictures_found = sum(len(slide.pictures) for slide in slides)
        captions = (
            await self._caption_all_pictures(slides)
            if self._captioner.enabled
            else None
        )

        text, boundaries = _assemble_text(slides)

        metadata: dict[str, Any] = {
            "slide_count": slide_count,
            SLIDE_BOUNDARIES_KEY: boundaries,
            "text_length": len(text),
            "parse_mode": "markdown",
            **picture_metadata(pictures_found, captions),
        }

        return ProcessingResult(
            text=text,
            metadata=metadata,
            processor=self.name,
            success=True,
        )

    async def health_check(self) -> bool:
        try:
            import pptx  # noqa: F401, PLC0415
        except ImportError:
            return False
        return True


def _render_table(table: Any) -> str:
    """A pptx table as a markdown table. Rows are always rectangular here,
    unlike a spreadsheet's sparse cell range, so no padding is needed."""
    return render_table(
        [[escape_cell(cell.text) for cell in row.cells] for row in table.rows]
    )


def _eligible_picture(shape: Any) -> Optional[Picture]:
    """``shape``'s image, if it is worth a docling captioning round trip."""
    try:
        image = shape.image
        return picture_if_eligible(image.content_type, image.size, image.blob)
    except Exception as exc:  # noqa: BLE001
        # A corrupt/unreadable image part costs this picture, not the deck.
        logger.debug("Skipping unreadable picture shape: %s", exc)
        return None


def _shape_blocks(shapes: Any) -> tuple[list[str], list[Picture]]:
    """Text frames, tables and eligible pictures in shape order, descending
    into group shapes."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: PLC0415 -- see _extract_deck

    blocks: list[str] = []
    pictures: list[Picture] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            sub_blocks, sub_pictures = _shape_blocks(shape.shapes)
            blocks.extend(sub_blocks)
            pictures.extend(sub_pictures)
        elif shape.has_table:
            blocks.append(_render_table(shape.table))
        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                blocks.append(text)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture = _eligible_picture(shape)
            if picture is not None:
                pictures.append(picture)
    return blocks, pictures


def _render_slide_markdown(slide: _SlideData) -> str:
    """One slide's blocks (text, tables, resolved image captions) plus notes."""
    blocks = list(slide.blocks)
    if slide.notes:
        blocks.append(f"**Notes:** {slide.notes}")
    if not blocks:
        return ""
    return f"## Slide {slide.index}\n\n" + "\n\n".join(blocks) + "\n"


def _assemble_text(slides: list[_SlideData]) -> tuple[str, list[dict[str, Any]]]:
    """Join every non-empty slide's markdown and its ``slide_boundaries``
    span, in one pass so the offsets always index the returned text exactly."""
    parts: list[str] = []
    boundaries: list[dict[str, Any]] = []
    offset = 0
    for slide in slides:
        body = _render_slide_markdown(slide)
        if not body:
            continue
        parts.append(body)
        boundaries.append(
            {
                "slide": slide.index,
                "start_offset": offset,
                "end_offset": offset + len(body),
            }
        )
        offset += len(body)
    return "".join(parts), boundaries


def _extract_deck(content: bytes) -> tuple[list[_SlideData], int]:
    """Parse every slide's blocks, notes and picture shapes. Runs in a worker
    thread; picture captioning is network I/O and happens afterwards, in the
    async ``process()``."""
    from pptx import Presentation  # noqa: PLC0415 -- keep the import off the hot path

    check_zip_size(content)

    prs = Presentation(io.BytesIO(content))

    slides: list[_SlideData] = []
    for i, slide in enumerate(prs.slides, start=1):
        blocks, pictures = _shape_blocks(slide.shapes)

        # A notes page need not carry a notes placeholder, in which case
        # notes_text_frame is None.
        notes_frame = (
            slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        )
        notes = notes_frame.text.strip() if notes_frame is not None else ""

        slides.append(
            _SlideData(index=i, blocks=blocks, pictures=pictures, notes=notes or None)
        )
    return slides, len(prs.slides)
